"""
Transaction management interface for spreadProfit.

Provides record entry, visual calendar selection, automatic unit conversions,
smart market price filling, CRUD operations, and multi-format data import/export
(Excel, CSV, JSON).
"""

from __future__ import annotations

import calendar
import csv
import datetime
import json
import os
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
from typing import Any, Callable, Dict, List, Optional

from config import (
    ASSET_CATEGORIES,
    CLR_BG,
    CLR_BORDER,
    CLR_ENTRY,
    CLR_GOLD,
    CLR_GREEN,
    CLR_HEAD,
    CLR_PANEL,
    CLR_RED,
    CLR_TEXT,
    CLR_TEXT_MUTED,
    DEFAULT_ASSETS,
    FONT_BOLD,
    FONT_MAIN,
    get_unit_factor,
    normalize_action,
    normalize_category,
    normalize_unit,
)
from db_manager import DATA_DIR, load_custom_assets, next_id, save_custom_assets, save_db


class CalendarDialog(tk.Toplevel):
    """Modal date selection calendar popup dialog."""

    def __init__(self, parent: tk.Widget, current_date_str: str, callback: Callable[[str], None]) -> None:
        super().__init__(parent)
        self.callback = callback
        self.title("Select Transaction Date")

        width, height = 440, 485
        self.configure(bg=CLR_PANEL)
        self.transient(parent)
        self.grab_set()
        self.resizable(False, False)

        # Center relative to parent
        self.update_idletasks()
        px, py = parent.winfo_rootx(), parent.winfo_rooty()
        pw, ph = parent.winfo_width(), parent.winfo_height()
        x = px + (pw - width) // 2
        y = py + (ph - height) // 2
        self.geometry(f"{width}x{height}+{x}+{y}")

        try:
            self.current_date = datetime.date.fromisoformat(current_date_str)
        except Exception:
            self.current_date = datetime.date.today()

        self.year = self.current_date.year
        self.month = self.current_date.month

        # Month / Year Navigation Header
        header = tk.Frame(self, bg=CLR_PANEL, pady=12)
        header.pack(fill="x")

        tk.Button(
            header,
            text="<",
            command=self.prev_month,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=("Segoe UI", 11, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
        ).pack(side="left", padx=15)

        self.lbl_month_year = tk.Label(
            header,
            text="",
            bg=CLR_PANEL,
            fg=CLR_GOLD,
            font=("Segoe UI", 13, "bold"),
        )
        self.lbl_month_year.pack(side="left", expand=True, fill="x")

        tk.Button(
            header,
            text=">",
            command=self.next_month,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=("Segoe UI", 11, "bold"),
            relief="flat",
            padx=10,
            pady=4,
            cursor="hand2",
        ).pack(side="right", padx=15)

        # Day of week header
        week_frame = tk.Frame(self, bg=CLR_PANEL)
        week_frame.pack(fill="x", padx=15)
        for i, day_name in enumerate(["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]):
            tk.Label(
                week_frame,
                text=day_name,
                bg=CLR_PANEL,
                fg=CLR_TEXT_MUTED,
                font=("Segoe UI", 11, "bold"),
                width=4,
            ).grid(row=0, column=i, padx=3, pady=4)

        # Calendar days grid
        self.days_frame = tk.Frame(self, bg=CLR_PANEL)
        self.days_frame.pack(fill="both", expand=True, padx=15, pady=8)
        for i in range(7):
            self.days_frame.columnconfigure(i, weight=1)

        self.draw_calendar()

    def draw_calendar(self) -> None:
        """Render month days onto the grid frame."""
        for child in self.days_frame.winfo_children():
            child.destroy()

        cal = calendar.Calendar(firstweekday=0)
        month_days = cal.monthdayscalendar(self.year, self.month)
        month_name = calendar.month_name[self.month]
        self.lbl_month_year.config(text=f"{month_name} {self.year}")

        today = datetime.date.today()
        for r_idx, week in enumerate(month_days):
            for c_idx, day in enumerate(week):
                if day == 0:
                    tk.Label(self.days_frame, text="", bg=CLR_PANEL, width=4, height=1).grid(
                        row=r_idx, column=c_idx, padx=3, pady=4
                    )
                else:
                    is_today = self.year == today.year and self.month == today.month and day == today.day
                    is_selected = (
                        self.year == self.current_date.year
                        and self.month == self.current_date.month
                        and day == self.current_date.day
                    )

                    bg_color = CLR_ENTRY
                    fg_color = CLR_TEXT
                    if is_selected:
                        bg_color = CLR_GOLD
                        fg_color = CLR_BG
                    elif is_today:
                        bg_color = CLR_HEAD
                        fg_color = "white"

                    btn = tk.Button(
                        self.days_frame,
                        text=str(day),
                        bg=bg_color,
                        fg=fg_color,
                        font=("Segoe UI", 11, "bold"),
                        relief="flat",
                        width=4,
                        height=1,
                        cursor="hand2",
                        command=lambda d=day: self.select_day(d),
                    )
                    btn.grid(row=r_idx, column=c_idx, padx=3, pady=4)

    def prev_month(self) -> None:
        """Navigate to preceding month."""
        self.month -= 1
        if self.month < 1:
            self.month = 12
            self.year -= 1
        self.draw_calendar()

    def next_month(self) -> None:
        """Navigate to subsequent month."""
        self.month += 1
        if self.month > 12:
            self.month = 1
            self.year += 1
        self.draw_calendar()

    def select_day(self, day: int) -> None:
        """Deliver selected date to callback and dismiss dialog."""
        selected_date = datetime.date(self.year, self.month, day)
        self.callback(selected_date.isoformat())
        self.destroy()


class TransactionsTab(tk.Frame):
    """Transaction management view supporting record editing, search, and data transfers."""

    def __init__(self, parent: tk.Widget, main_app: Any) -> None:
        super().__init__(parent, bg=CLR_BG)
        self.main_app = main_app

        today = datetime.date.today()
        self.var_day = tk.StringVar(value=f"{today.day:02d}")
        self.var_month = tk.StringVar(value=f"{today.month:02d}")
        self.var_year = tk.StringVar(value=str(today.year))
        self.var_category = tk.StringVar(value=ASSET_CATEGORIES[0])
        self.var_asset_select = tk.StringVar()
        self.var_custom_name = tk.StringVar()
        self.var_action = tk.StringVar(value="Buy")
        self.var_qty = tk.StringVar()
        self.var_unit = tk.StringVar(value="Ounce")
        self.var_unit_price = tk.StringVar()
        self.var_note = tk.StringVar()
        self.var_search = tk.StringVar()
        self.var_filter_cat = tk.StringVar(value="All")
        self.var_total_preview = tk.StringVar(value="Estimated Total: -")
        self.editing_id: Optional[int] = None
        self.last_unit: str = "Ounce"

        self._build_ui()
        self._on_category_change()

        # Update preview whenever quantity or price changes
        self.var_qty.trace_add("write", self._update_total_preview)
        self.var_unit_price.trace_add("write", self._update_total_preview)

    def _build_ui(self) -> None:
        """Assemble split panes for transaction entry form and historical record table."""
        paned = tk.PanedWindow(self, orient="horizontal", bg=CLR_BG, sashwidth=6, sashrelief="flat")
        paned.pack(fill="both", expand=True, padx=8, pady=8)

        left_panel = tk.Frame(paned, bg=CLR_BG)
        right_panel = tk.Frame(paned, bg=CLR_BG)
        paned.add(left_panel, minsize=420)
        paned.add(right_panel, minsize=650)

        # ---------------- Left Panel: Transaction Entry Form ----------------
        form_frame = tk.LabelFrame(
            left_panel,
            text=" TRANSACTION ENTRY ",
            font=FONT_BOLD,
            bg=CLR_PANEL,
            fg=CLR_GOLD,
            bd=1,
            relief="solid",
            padx=10,
            pady=8,
        )
        form_frame.pack(fill="both", expand=True)

        def make_label(text: str, row: int, col: int) -> None:
            tk.Label(form_frame, text=text, bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).grid(
                row=row, column=col, sticky="e", padx=5, pady=4
            )

        # 1. Date inputs
        make_label("Date (D/M/Y):", 0, 0)
        date_frame = tk.Frame(form_frame, bg=CLR_PANEL)
        date_frame.grid(row=0, column=1, sticky="ew", padx=5, pady=4)

        self.ent_day = tk.Entry(
            date_frame,
            textvariable=self.var_day,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
            width=4,
            justify="center",
        )
        self.ent_day.pack(side="left")

        tk.Label(date_frame, text="/", bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).pack(side="left", padx=2)
        self.ent_month = tk.Entry(
            date_frame,
            textvariable=self.var_month,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
            width=4,
            justify="center",
        )
        self.ent_month.pack(side="left")

        tk.Label(date_frame, text="/", bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).pack(side="left", padx=2)
        self.ent_year = tk.Entry(
            date_frame,
            textvariable=self.var_year,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
            width=6,
            justify="center",
        )
        self.ent_year.pack(side="left")

        tk.Button(
            date_frame,
            text="Cal",
            command=self._show_calendar,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            padx=6,
            pady=2,
            cursor="hand2",
        ).pack(side="left", padx=(10, 0))

        # Auto-jump key bindings
        self.ent_day.bind("<Key>", self._on_day_key)
        self.ent_month.bind("<Key>", self._on_month_key)
        self.ent_month.bind("<BackSpace>", self._on_month_backspace)
        self.ent_year.bind("<BackSpace>", self._on_year_backspace)

        # 2. Asset Category
        make_label("Category:", 1, 0)
        cb_cat = ttk.Combobox(
            form_frame,
            textvariable=self.var_category,
            values=ASSET_CATEGORIES,
            state="readonly",
            font=FONT_MAIN,
        )
        cb_cat.grid(row=1, column=1, sticky="ew", padx=5, pady=4)
        cb_cat.bind("<<ComboboxSelected>>", self._on_category_change)

        # 3. Asset Selector
        make_label("Specific Asset:", 2, 0)
        asset_select_frame = tk.Frame(form_frame, bg=CLR_PANEL)
        asset_select_frame.grid(row=2, column=1, sticky="ew", padx=5, pady=4)
        asset_select_frame.columnconfigure(0, weight=1)

        self.cb_asset = ttk.Combobox(
            asset_select_frame,
            textvariable=self.var_asset_select,
            state="readonly",
            font=FONT_MAIN,
        )
        self.cb_asset.grid(row=0, column=0, sticky="ew", padx=(0, 5))
        self.cb_asset.bind("<<ComboboxSelected>>", self._on_asset_select_change)

        self.btn_edit_asset = tk.Button(
            asset_select_frame,
            text="Edit",
            command=self._edit_custom_asset,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=("Segoe UI", 9),
            relief="flat",
            padx=5,
            cursor="hand2",
            state="disabled",
        )
        self.btn_edit_asset.grid(row=0, column=1, sticky="e", padx=2)

        self.btn_del_asset = tk.Button(
            asset_select_frame,
            text="Del",
            command=self._delete_custom_asset,
            bg=CLR_ENTRY,
            fg=CLR_RED,
            font=("Segoe UI", 9),
            relief="flat",
            padx=5,
            cursor="hand2",
            state="disabled",
        )
        self.btn_del_asset.grid(row=0, column=2, sticky="e", padx=2)

        # 4. Custom Asset Creation Panel (Visible when 'Custom / Other' selected)
        self.custom_assets_frame = tk.Frame(form_frame, bg=CLR_PANEL)
        self.custom_assets_frame.grid(row=3, column=0, columnspan=2, sticky="ew")

        tk.Label(
            self.custom_assets_frame,
            text="New Asset Name:",
            bg=CLR_PANEL,
            fg=CLR_TEXT,
            font=FONT_BOLD,
        ).grid(row=0, column=0, sticky="e", padx=5, pady=4)

        cust_name_box = tk.Frame(self.custom_assets_frame, bg=CLR_PANEL)
        cust_name_box.grid(row=0, column=1, sticky="ew", padx=5, pady=4)
        cust_name_box.columnconfigure(0, weight=1)

        self.ent_cust_name = tk.Entry(
            cust_name_box,
            textvariable=self.var_custom_name,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
        )
        self.ent_cust_name.grid(row=0, column=0, sticky="ew", padx=(0, 5))

        tk.Button(
            cust_name_box,
            text="Save Asset",
            command=self._save_new_custom_asset,
            bg="#4f46e5",
            fg="white",
            font=("Segoe UI", 9, "bold"),
            relief="flat",
            padx=8,
            cursor="hand2",
        ).grid(row=0, column=1, sticky="e")
        self.custom_assets_frame.columnconfigure(1, weight=1)

        # 5. Transaction Action
        make_label("Action:", 4, 0)
        cb_act = ttk.Combobox(
            form_frame,
            textvariable=self.var_action,
            values=["Buy", "Sell"],
            state="readonly",
            font=FONT_MAIN,
        )
        cb_act.grid(row=4, column=1, sticky="ew", padx=5, pady=4)
        cb_act.bind("<<ComboboxSelected>>", lambda _: self._update_price_from_market())

        # 6. Quantity & Unit
        make_label("Quantity:", 5, 0)
        qty_unit_box = tk.Frame(form_frame, bg=CLR_PANEL)
        qty_unit_box.grid(row=5, column=1, sticky="ew")
        qty_unit_box.columnconfigure(0, weight=2)
        qty_unit_box.columnconfigure(2, weight=1)

        tk.Entry(
            qty_unit_box,
            textvariable=self.var_qty,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
        ).grid(row=0, column=0, sticky="ew", padx=(0, 5))

        tk.Label(qty_unit_box, text="Unit:", bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).grid(
            row=0, column=1, sticky="e", padx=2
        )
        self.cb_unit = ttk.Combobox(
            qty_unit_box,
            textvariable=self.var_unit,
            values=[
                "Ounce", "Gram", "Kilogram", "Tael", "Chi",
                "0.1 Chi", "0.2 Chi", "0.5 Chi", "1 Chi", "2 Chi", "5 Chi",
            ],
            font=FONT_MAIN,
            width=8,
        )
        self.cb_unit.grid(row=0, column=2, sticky="ew")
        self.cb_unit.bind("<<ComboboxSelected>>", self._on_unit_selected)

        # 7. Unit Price
        make_label("Unit Price (VND):", 6, 0)
        price_box = tk.Frame(form_frame, bg=CLR_PANEL)
        price_box.grid(row=6, column=1, sticky="ew")
        price_box.columnconfigure(0, weight=1)

        self.ent_price = tk.Entry(
            price_box,
            textvariable=self.var_unit_price,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
        )
        self.ent_price.grid(row=0, column=0, sticky="ew", padx=(0, 5))
        self.ent_price.bind("<FocusOut>", self._format_price_entry)
        self.ent_price.bind("<FocusIn>", self._strip_price_entry)
        self.ent_price.bind("<Return>", self._format_price_entry)

        tk.Button(
            price_box,
            text="Use Benchmark",
            command=self._update_price_from_market,
            bg="#3b82f6",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            padx=6,
            pady=2,
            cursor="hand2",
        ).grid(row=0, column=1, sticky="e")

        # 8. Note
        make_label("Note:", 7, 0)
        tk.Entry(
            form_frame,
            textvariable=self.var_note,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            insertbackground=CLR_TEXT,
            bd=3,
        ).grid(row=7, column=1, sticky="ew", padx=5, pady=4)

        # 9. Estimated Total Preview
        tk.Label(
            form_frame,
            textvariable=self.var_total_preview,
            bg=CLR_PANEL,
            fg=CLR_GOLD,
            font=FONT_BOLD,
            anchor="w",
        ).grid(row=8, column=0, columnspan=2, sticky="w", padx=10, pady=10)

        form_frame.columnconfigure(1, weight=1)

        # Action Buttons
        btn_box = tk.Frame(form_frame, bg=CLR_PANEL)
        btn_box.grid(row=9, column=0, columnspan=2, sticky="ew", pady=(10, 0))

        self.btn_add = tk.Button(
            btn_box,
            text="Add Transaction",
            command=self._add_transaction,
            bg=CLR_GREEN,
            fg="white",
            font=FONT_BOLD,
            relief="flat",
            padx=10,
            pady=6,
            cursor="hand2",
        )
        self.btn_add.pack(side="left", padx=4)

        self.btn_cancel = tk.Button(
            btn_box,
            text="Cancel",
            command=self._cancel_edit,
            bg="#64748b",
            fg="white",
            font=FONT_BOLD,
            relief="flat",
            padx=10,
            pady=6,
            cursor="hand2",
        )

        self.btn_load = tk.Button(
            btn_box,
            text="Edit Row",
            command=self._load_selected,
            bg="#555588",
            fg="white",
            font=FONT_BOLD,
            relief="flat",
            padx=10,
            pady=6,
            cursor="hand2",
        )
        self.btn_load.pack(side="left", padx=4)

        self.btn_del = tk.Button(
            btn_box,
            text="Delete",
            command=self._delete_selected,
            bg=CLR_RED,
            fg="white",
            font=FONT_BOLD,
            relief="flat",
            padx=10,
            pady=6,
            cursor="hand2",
        )
        self.btn_del.pack(side="left", padx=4)

        # ---------------- Right Panel: Transaction History ----------------
        hist_frame = tk.LabelFrame(
            right_panel,
            text=" TRANSACTION HISTORY ",
            font=FONT_BOLD,
            bg=CLR_PANEL,
            fg=CLR_GOLD,
            bd=1,
            relief="solid",
            padx=10,
            pady=8,
        )
        hist_frame.pack(fill="both", expand=True)

        filter_bar = tk.Frame(hist_frame, bg=CLR_PANEL)
        filter_bar.pack(fill="x", pady=(0, 10))

        tk.Label(filter_bar, text="Search:", bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).pack(side="left", padx=2)
        ent_search = tk.Entry(
            filter_bar,
            textvariable=self.var_search,
            bg=CLR_ENTRY,
            fg=CLR_TEXT,
            font=FONT_MAIN,
            relief="flat",
            width=15,
            bd=3,
        )
        ent_search.pack(side="left", padx=5)
        self.var_search.trace_add("write", lambda *_: self.refresh_table())

        tk.Label(filter_bar, text="Category:", bg=CLR_PANEL, fg=CLR_TEXT, font=FONT_BOLD).pack(side="left", padx=5)
        cb_fcat = ttk.Combobox(
            filter_bar,
            textvariable=self.var_filter_cat,
            values=["All"] + ASSET_CATEGORIES,
            state="readonly",
            font=FONT_MAIN,
            width=12,
        )
        cb_fcat.pack(side="left", padx=5)
        cb_fcat.bind("<<ComboboxSelected>>", lambda _: self.refresh_table())

        # Data export / import action buttons
        tk.Button(
            filter_bar,
            text="Export Data",
            command=self._export_data,
            bg=CLR_BORDER,
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            padx=10,
            pady=3,
            cursor="hand2",
        ).pack(side="right", padx=4)

        tk.Button(
            filter_bar,
            text="Import Data",
            command=self._import_data,
            bg="#4f46e5",
            fg="white",
            font=("Segoe UI", 8, "bold"),
            relief="flat",
            padx=10,
            pady=3,
            cursor="hand2",
        ).pack(side="right", padx=4)

        # Main Records Treeview
        columns = ("ID", "Date", "Category", "Asset", "Action", "Qty", "Unit", "Price", "Total", "Note")
        self.tree = ttk.Treeview(hist_frame, columns=columns, show="headings", height=18, selectmode="browse")

        self.tree.heading("ID", text="ID")
        self.tree.heading("Date", text="Date")
        self.tree.heading("Category", text="Category")
        self.tree.heading("Asset", text="Asset Name")
        self.tree.heading("Action", text="Action")
        self.tree.heading("Qty", text="Qty")
        self.tree.heading("Unit", text="Unit")
        self.tree.heading("Price", text="Unit Price")
        self.tree.heading("Total", text="Total Amount")
        self.tree.heading("Note", text="Note")

        self.tree.column("ID", width=35, anchor="center")
        self.tree.column("Date", width=85, anchor="center")
        self.tree.column("Category", width=80, anchor="center")
        self.tree.column("Asset", width=200, anchor="w")
        self.tree.column("Action", width=55, anchor="center")
        self.tree.column("Qty", width=55, anchor="center")
        self.tree.column("Unit", width=45, anchor="center")
        self.tree.column("Price", width=105, anchor="e")
        self.tree.column("Total", width=115, anchor="e")
        self.tree.column("Note", width=160, anchor="w")

        vsb = ttk.Scrollbar(hist_frame, orient="vertical", command=self.tree.yview)
        hsb = ttk.Scrollbar(hist_frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vsb.set, xscrollcommand=hsb.set)

        self.tree.pack(side="top", fill="both", expand=True)
        vsb.pack(side="right", fill="y", before=self.tree)
        hsb.pack(side="bottom", fill="x")

        self.tree.tag_configure("buy", background="#132c20", foreground="#6ee7b7")
        self.tree.tag_configure("sell", background="#3b1919", foreground="#fca5a5")

    def _on_category_change(self, event: Any = None) -> None:
        """Update available assets in dropdown when switching categories."""
        cat = normalize_category(self.var_category.get())
        assets = [a for a in DEFAULT_ASSETS if a["category"] == cat]
        custom_assets = load_custom_assets()
        assets += [a for a in custom_assets if a.get("category") == cat]

        values = [a["name"] for a in assets] + ["Custom / Other..."]
        self.cb_asset.configure(values=values)
        if values:
            self.cb_asset.current(0)
            self._on_asset_select_change()

    def _on_asset_select_change(self, event: Any = None) -> None:
        """Handle asset selection change, revealing custom inputs if selected."""
        val = self.var_asset_select.get()
        custom_assets = load_custom_assets()
        is_custom = any(a["name"] == val for a in custom_assets)

        state = "normal" if is_custom else "disabled"
        self.btn_edit_asset.configure(state=state)
        self.btn_del_asset.configure(state=state)

        if val == "Custom / Other...":
            self.custom_assets_frame.grid()
            self.var_custom_name.set("")
            self.var_unit.set("Tael")
        else:
            self.custom_assets_frame.grid_remove()
            asset_info = next((a for a in DEFAULT_ASSETS if a["name"] == val), None)
            if not asset_info:
                asset_info = next((a for a in custom_assets if a["name"] == val), None)
            if asset_info:
                self.var_custom_name.set(asset_info["name"])
                self.var_unit.set(normalize_unit(asset_info["unit"]))

        self.last_unit = self.var_unit.get()
        self._update_price_from_market()

    def _on_unit_selected(self, event: Any = None) -> None:
        """Automatically recalculate quantity and unit price when switching units."""
        new_unit = normalize_unit(self.var_unit.get().strip())
        old_unit = getattr(self, "last_unit", "")
        self.last_unit = new_unit

        if not old_unit or old_unit == new_unit:
            return

        old_factor = get_unit_factor(old_unit)
        new_factor = get_unit_factor(new_unit)

        if old_factor > 0 and new_factor > 0:
            try:
                qty_val = float(self.var_qty.get())
                price_val = float(self.var_unit_price.get().replace(",", ""))
                ratio = old_factor / new_factor
                new_qty = qty_val * ratio
                new_price = price_val / ratio

                self.var_qty.set(f"{new_qty:g}")
                self.var_unit_price.set(f"{new_price:,.0f}")
                self.main_app.var_status.set(f"Unit converted: {qty_val:g} {old_unit} -> {new_qty:g} {new_unit}")
            except ValueError:
                pass

    def _update_total_preview(self, *_: Any) -> None:
        """Recalculate gross transaction total preview."""
        try:
            qty = float(self.var_qty.get())
            price = float(self.var_unit_price.get().replace(",", ""))
            total = qty * price
            self.var_total_preview.set(f"Estimated Total: {total:,.0f} VND")
        except ValueError:
            self.var_total_preview.set("Estimated Total: -")

    def get_date(self) -> str:
        """Assemble ISO-8601 YYYY-MM-DD date string from day, month, and year inputs."""
        try:
            d_str = self.var_day.get().strip()
            m_str = self.var_month.get().strip()
            y_str = self.var_year.get().strip()
            if not d_str or not m_str or not y_str:
                return ""
            d, m, y = int(d_str), int(m_str), int(y_str)
            datetime.date(y, m, d)
            return f"{y:04d}-{m:02d}-{d:02d}"
        except ValueError:
            return ""

    def set_date(self, date_str: str) -> None:
        """Decompose ISO date string into day, month, and year entry values."""
        try:
            d = datetime.date.fromisoformat(date_str)
            self.var_day.set(f"{d.day:02d}")
            self.var_month.set(f"{d.month:02d}")
            self.var_year.set(str(d.year))
        except Exception:
            pass

    def _format_price_entry(self, event: Any = None) -> None:
        """Format price string with comma thousands grouping."""
        try:
            val = self.var_unit_price.get()
            clean = "".join([c for c in val.replace(",", "") if c.isdigit()])
            if clean:
                self.var_unit_price.set(f"{int(clean):,}")
        except Exception:
            pass

    def _strip_price_entry(self, event: Any = None) -> None:
        """Strip comma formatting during text entry."""
        try:
            clean = self.var_unit_price.get().replace(",", "")
            self.var_unit_price.set(clean)
            self.ent_price.icursor(tk.END)
        except Exception:
            pass

    def _show_calendar(self) -> None:
        """Display date selector dialog."""
        CalendarDialog(self, self.get_date(), self.set_date)

    def _update_price_from_market(self) -> None:
        """Autofill unit price using current benchmark market quotes."""
        cat = normalize_category(self.var_category.get())
        price_info = self.main_app.current_prices.get(cat)
        if not price_info:
            legacy_key = "Vàng" if cat == "Gold" else "Bạc"
            price_info = self.main_app.current_prices.get(legacy_key)

        if price_info:
            action = normalize_action(self.var_action.get())
            price = price_info["sell"] if action == "Buy" else price_info["buy"]
            price_unit = normalize_unit(price_info.get("unit", "Tael"))
            form_unit = normalize_unit(self.var_unit.get().strip())

            if price > 0:
                ratio = get_unit_factor(price_unit) / get_unit_factor(form_unit)
                converted_price = price / ratio if ratio > 0 else price
                self.var_unit_price.set(f"{converted_price:,.0f}")
                self.main_app.var_status.set(
                    f"Benchmark applied for {cat} ({action}): {converted_price:,.0f} VND / {form_unit}"
                )

    def refresh_table(self) -> None:
        """Filter and display historical transaction entries."""
        for row in self.tree.get_children():
            self.tree.delete(row)

        search_kw = self.var_search.get().lower().strip()
        filter_cat = self.var_filter_cat.get()

        for r in sorted(self.main_app.db, key=lambda x: str(x.get("date", "")), reverse=True):
            cat = normalize_category(r.get("category", "Gold"))
            if filter_cat != "All" and cat != filter_cat:
                continue

            asset_name = r.get("asset", "").lower()
            note = r.get("note", "").lower()
            date_str = str(r.get("date", ""))

            if search_kw and (search_kw not in asset_name and search_kw not in note and search_kw not in date_str):
                continue

            action = normalize_action(r.get("action", "Buy"))
            tag = "buy" if action == "Buy" else "sell"

            self.tree.insert(
                "",
                "end",
                tags=(tag,),
                values=(
                    r["id"],
                    date_str,
                    cat,
                    r["asset"],
                    action,
                    f"{r['quantity']:g}",
                    normalize_unit(r["unit"]),
                    f"{r['unit_price']:,.0f}",
                    f"{r['total']:,.0f}",
                    r.get("note", ""),
                ),
            )

    def _add_transaction(self) -> None:
        """Validate input fields and commit a new or edited transaction record."""
        try:
            qty = float(self.var_qty.get())
            price = float(self.var_unit_price.get().replace(",", ""))
        except ValueError:
            messagebox.showerror("Validation Error", "Quantity and unit price must be valid numbers.")
            return

        asset_select = self.var_asset_select.get()
        if asset_select == "Custom / Other...":
            asset_name = self.var_custom_name.get().strip()
            if not asset_name:
                messagebox.showerror("Validation Error", "Asset name is required.")
                return

            custom_assets = load_custom_assets()
            if not any(a["name"].lower() == asset_name.lower() for a in DEFAULT_ASSETS) and not any(
                a["name"].lower() == asset_name.lower() for a in custom_assets
            ):
                new_asset = {
                    "name": asset_name,
                    "category": normalize_category(self.var_category.get()),
                    "unit": normalize_unit(self.var_unit.get().strip()),
                }
                custom_assets.append(new_asset)
                save_custom_assets(custom_assets)
        else:
            asset_name = asset_select

        tx_date = self.get_date()
        if not tx_date:
            messagebox.showerror("Validation Error", "Please provide a valid transaction date.")
            return

        cat = normalize_category(self.var_category.get())
        act = normalize_action(self.var_action.get())
        unit = normalize_unit(self.var_unit.get().strip())

        if self.editing_id is not None:
            # Edit existing record
            rec = next((r for r in self.main_app.db if r["id"] == self.editing_id), None)
            if not rec:
                messagebox.showerror("Error", "Record not found.")
                self._cancel_edit()
                return

            rec["date"] = tx_date
            rec["category"] = cat
            rec["asset"] = asset_name
            rec["action"] = act
            rec["quantity"] = qty
            rec["unit"] = unit
            rec["unit_price"] = price
            rec["total"] = qty * price
            rec["note"] = self.var_note.get().strip()

            save_db(self.main_app.db)
            self.editing_id = None
            self.btn_add.configure(text="Add Transaction", bg=CLR_GREEN)
            self.btn_cancel.pack_forget()
            msg = f"Updated transaction ID {rec['id']}."
        else:
            # Create fresh record
            rec = {
                "id": next_id(self.main_app.db),
                "date": tx_date,
                "category": cat,
                "asset": asset_name,
                "action": act,
                "quantity": qty,
                "unit": unit,
                "unit_price": price,
                "total": qty * price,
                "note": self.var_note.get().strip(),
            }
            self.main_app.db.append(rec)
            save_db(self.main_app.db)
            msg = f"Recorded {rec['action']} {qty:g} {rec['unit']} of {rec['asset']}."

        # Reset form fields
        self.var_qty.set("")
        self.var_unit_price.set("")
        self.var_note.set("")
        self.var_total_preview.set("Estimated Total: -")

        self._on_category_change()
        self.refresh_table()
        self.main_app.refresh_portfolio_view()
        self.main_app.var_status.set(msg)

    def _cancel_edit(self) -> None:
        """Reset edit state back to creation mode."""
        self.editing_id = None
        self.btn_add.configure(text="Add Transaction", bg=CLR_GREEN)
        self.btn_cancel.pack_forget()
        self.var_qty.set("")
        self.var_unit_price.set("")
        self.var_note.set("")
        self.var_total_preview.set("Estimated Total: -")

    def _load_selected(self) -> None:
        """Load selected row data into entry form for editing."""
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Notice", "Please select a transaction to edit.")
            return

        iid = int(self.tree.item(sel[0])["values"][0])
        rec = next((r for r in self.main_app.db if r["id"] == iid), None)
        if not rec:
            return

        self.editing_id = iid
        self.btn_add.configure(text="Save Changes", bg=CLR_GOLD)

        # Repack buttons in proper order
        for btn in (self.btn_add, self.btn_cancel, self.btn_load, self.btn_del):
            btn.pack_forget()
        for btn in (self.btn_add, self.btn_cancel, self.btn_load, self.btn_del):
            btn.pack(side="left", padx=4)

        self.set_date(rec.get("date", ""))
        self.var_category.set(normalize_category(rec.get("category", ASSET_CATEGORIES[0])))
        self._on_category_change()

        found_in_default = any(a["name"] == rec["asset"] for a in DEFAULT_ASSETS)
        if found_in_default:
            self.var_asset_select.set(rec["asset"])
            self.custom_assets_frame.grid_remove()
        else:
            self.var_asset_select.set("Custom / Other...")
            self.custom_assets_frame.grid()
            self.var_custom_name.set(rec["asset"])

        self.var_action.set(normalize_action(rec["action"]))
        self.var_qty.set(f"{rec['quantity']:g}")
        self.var_unit.set(normalize_unit(rec["unit"]))
        self.var_unit_price.set(f"{rec['unit_price']:,.0f}")
        self.last_unit = normalize_unit(rec["unit"])
        self.var_note.set(rec.get("note", ""))

    def _delete_selected(self) -> None:
        """Remove selected transaction record after confirmation."""
        sel = self.tree.selection()
        if not sel:
            messagebox.showwarning("Notice", "Please select a row to delete.")
            return

        iid = int(self.tree.item(sel[0])["values"][0])
        if not messagebox.askyesno("Confirm Deletion", f"Permanently delete transaction ID {iid}?"):
            return

        self.main_app.db = [r for r in self.main_app.db if r["id"] != iid]
        save_db(self.main_app.db)
        self.refresh_table()
        self.main_app.refresh_portfolio_view()
        self.main_app.var_status.set(f"Deleted transaction ID {iid}.")

    def _edit_custom_asset(self) -> None:
        """Rename selected custom asset in catalog and past transaction history."""
        val = self.var_asset_select.get()
        custom_assets = load_custom_assets()
        asset_info = next((a for a in custom_assets if a["name"] == val), None)
        if not asset_info:
            return

        new_name = simpledialog.askstring("Rename Asset", f"Enter new name for '{val}':", initialvalue=val)
        if not new_name or not new_name.strip():
            return
        new_name = new_name.strip()

        for a in custom_assets:
            if a["name"] == val:
                a["name"] = new_name

        for tx in self.main_app.db:
            if tx.get("asset") == val:
                tx["asset"] = new_name

        save_custom_assets(custom_assets)
        save_db(self.main_app.db)

        self._on_category_change()
        self.var_asset_select.set(new_name)
        self._on_asset_select_change()
        self.refresh_table()
        self.main_app.var_status.set(f"Renamed asset from '{val}' to '{new_name}'.")

    def _delete_custom_asset(self) -> None:
        """Remove custom asset from repository."""
        val = self.var_asset_select.get()
        custom_assets = load_custom_assets()
        if not any(a["name"] == val for a in custom_assets):
            return

        if messagebox.askyesno("Confirm", f"Remove custom asset '{val}' from catalog?"):
            custom_assets = [a for a in custom_assets if a["name"] != val]
            save_custom_assets(custom_assets)
            self._on_category_change()
            self.main_app.var_status.set(f"Removed custom asset '{val}'.")

    def _save_new_custom_asset(self) -> None:
        """Register newly entered custom asset."""
        asset_name = self.var_custom_name.get().strip()
        if not asset_name:
            messagebox.showerror("Error", "Asset name cannot be blank.")
            return

        custom_assets = load_custom_assets()
        if any(a["name"].lower() == asset_name.lower() for a in DEFAULT_ASSETS) or any(
            a["name"].lower() == asset_name.lower() for a in custom_assets
        ):
            messagebox.showerror("Error", "An asset with this name already exists.")
            return

        new_asset = {
            "name": asset_name,
            "category": normalize_category(self.var_category.get()),
            "unit": normalize_unit(self.var_unit.get().strip() or "Tael"),
        }
        custom_assets.append(new_asset)
        save_custom_assets(custom_assets)

        self._on_category_change()
        self.var_asset_select.set(asset_name)
        self._on_asset_select_change()
        self.main_app.var_status.set(f"Registered new asset '{asset_name}'.")

    def _on_day_key(self, event: Any) -> Optional[str]:
        """Auto-jump to Month entry upon separator or 2 digits."""
        if event.char in ("/", ".", " ", "-"):
            self.ent_month.focus_set()
            self.ent_month.selection_range(0, tk.END)
            return "break"
        self.after(10, self._check_day_jump)
        return None

    def _check_day_jump(self) -> None:
        val = "".join([c for c in self.var_day.get() if c.isdigit()])
        if len(val) >= 2:
            self.ent_month.focus_set()
            self.ent_month.selection_range(0, tk.END)

    def _on_month_key(self, event: Any) -> Optional[str]:
        """Auto-jump to Year entry upon separator or 2 digits."""
        if event.char in ("/", ".", " ", "-"):
            self.ent_year.focus_set()
            self.ent_year.selection_range(0, tk.END)
            return "break"
        self.after(10, self._check_month_jump)
        return None

    def _check_month_jump(self) -> None:
        val = "".join([c for c in self.var_month.get() if c.isdigit()])
        if len(val) >= 2:
            self.ent_year.focus_set()
            self.ent_year.selection_range(0, tk.END)

    def _on_month_backspace(self, event: Any) -> Optional[str]:
        """Backspace jump back to Day input."""
        if not self.var_month.get():
            self.ent_day.focus_set()
            val = self.var_day.get()
            if val:
                self.var_day.set(val[:-1])
            self.ent_day.icursor(tk.END)
            return "break"
        return None

    def _on_year_backspace(self, event: Any) -> Optional[str]:
        """Backspace jump back to Month input."""
        if not self.var_year.get():
            self.ent_month.focus_set()
            val = self.var_month.get()
            if val:
                self.var_month.set(val[:-1])
            self.ent_month.icursor(tk.END)
            return "break"
        return None

    def _export_data(self) -> None:
        """Export transaction records to Excel (.xlsx), CSV, or JSON format."""
        if not self.main_app.db:
            messagebox.showwarning("Notice", "No transactions available to export.")
            return

        exports_dir = os.path.join(DATA_DIR, "exports")
        os.makedirs(exports_dir, exist_ok=True)
        default_name = f"spreadprofit_transactions_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}"

        filepath = filedialog.asksaveasfilename(
            title="Export Transaction Records",
            initialdir=exports_dir,
            initialfile=default_name,
            defaultextension=".xlsx",
            filetypes=[
                ("Excel Spreadsheet", "*.xlsx"),
                ("CSV Comma Separated", "*.csv"),
                ("JSON Database", "*.json"),
            ],
        )
        if not filepath:
            return

        try:
            if filepath.endswith(".json"):
                with open(filepath, "w", encoding="utf-8") as f:
                    json.dump(self.main_app.db, f, ensure_ascii=False, indent=2)
            elif filepath.endswith(".csv"):
                with open(filepath, "w", encoding="utf-8", newline="") as f:
                    writer = csv.writer(f)
                    writer.writerow(["ID", "Date", "Category", "Asset", "Action", "Quantity", "Unit", "UnitPrice", "Total", "Note"])
                    for r in self.main_app.db:
                        writer.writerow([
                            r["id"],
                            r.get("date", ""),
                            normalize_category(r.get("category", "")),
                            r["asset"],
                            normalize_action(r["action"]),
                            r["quantity"],
                            normalize_unit(r["unit"]),
                            r["unit_price"],
                            r["total"],
                            r.get("note", ""),
                        ])
            elif filepath.endswith(".xlsx"):
                import openpyxl

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Transactions"
                ws.append(["ID", "Date", "Category", "Asset", "Action", "Quantity", "Unit", "UnitPrice", "Total", "Note"])
                for r in self.main_app.db:
                    ws.append([
                        r["id"],
                        r.get("date", ""),
                        normalize_category(r.get("category", "")),
                        r["asset"],
                        normalize_action(r["action"]),
                        r["quantity"],
                        normalize_unit(r["unit"]),
                        r["unit_price"],
                        r["total"],
                        r.get("note", ""),
                    ])
                wb.save(filepath)

            messagebox.showinfo("Success", f"Data exported successfully to:\n{filepath}")
            self.main_app.var_status.set("Export complete.")
        except Exception as e:
            messagebox.showerror("Export Error", f"Failed to export data: {e}")

    def _import_data(self) -> None:
        """Import transaction records from Excel, CSV, or JSON with merge/replace option."""
        exports_dir = os.path.join(DATA_DIR, "exports")
        filepath = filedialog.askopenfilename(
            title="Import Transaction Records",
            initialdir=exports_dir,
            filetypes=[
                ("All Supported Files", "*.json;*.csv;*.xlsx;*.xls"),
                ("Excel Files", "*.xlsx;*.xls"),
                ("CSV Files", "*.csv"),
                ("JSON Files", "*.json"),
            ],
        )
        if not filepath:
            return

        try:
            imported_records: List[Dict[str, Any]] = []

            if filepath.endswith(".json"):
                with open(filepath, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if not isinstance(data, list):
                    raise ValueError("JSON file must contain an array of transaction objects.")
                for item in data:
                    qty = float(item["quantity"])
                    price = float(item["unit_price"])
                    imported_records.append({
                        "date": str(item["date"]),
                        "category": normalize_category(str(item.get("category", "Gold"))),
                        "asset": str(item["asset"]),
                        "action": normalize_action(str(item["action"])),
                        "quantity": qty,
                        "unit": normalize_unit(str(item["unit"])),
                        "unit_price": price,
                        "total": float(item.get("total", qty * price)),
                        "note": str(item.get("note", "")),
                    })

            elif filepath.endswith(".csv"):
                with open(filepath, "r", encoding="utf-8") as f:
                    reader = csv.reader(f)
                    header = next(reader, None)
                    if not header:
                        raise ValueError("CSV file is empty.")
                    norm_headers = [h.strip().lower() for h in header]

                    def find_col(aliases: List[str]) -> int:
                        for alias in aliases:
                            if alias in norm_headers:
                                return norm_headers.index(alias)
                        return -1

                    idx_date = find_col(["date", "ngày", "ngay"])
                    idx_cat = find_col(["category", "nhóm", "nhom"])
                    idx_asset = find_col(["asset", "tên tài sản", "ten tai san", "tên"])
                    idx_action = find_col(["action", "hành động", "hanh dong", "gd"])
                    idx_qty = find_col(["quantity", "số lượng", "so luong", "sl"])
                    idx_unit = find_col(["unit", "đơn vị", "don vi", "đv"])
                    idx_price = find_col(["unitprice", "unit_price", "đơn giá", "don gia", "giá"])
                    idx_total = find_col(["total", "tổng tiền", "tong tien", "tổng"])
                    idx_note = find_col(["note", "ghi chú", "ghi chu"])

                    for row in reader:
                        if not row or len(row) <= max(idx_date, idx_asset, idx_qty, idx_price):
                            continue
                        qty = float(row[idx_qty].replace(",", "").strip())
                        price = float(row[idx_price].replace(",", "").strip())
                        total = (
                            float(row[idx_total].replace(",", "").strip())
                            if idx_total != -1 and row[idx_total].strip()
                            else qty * price
                        )
                        imported_records.append({
                            "date": row[idx_date].strip(),
                            "category": normalize_category(row[idx_cat].strip() if idx_cat != -1 else "Gold"),
                            "asset": row[idx_asset].strip(),
                            "action": normalize_action(row[idx_action].strip() if idx_action != -1 else "Buy"),
                            "quantity": qty,
                            "unit": normalize_unit(row[idx_unit].strip() if idx_unit != -1 else "Tael"),
                            "unit_price": price,
                            "total": total,
                            "note": row[idx_note].strip() if idx_note != -1 else "",
                        })

            elif filepath.endswith((".xlsx", ".xls")):
                import openpyxl

                wb = openpyxl.load_workbook(filepath, data_only=True)
                ws = wb.active
                rows = list(ws.iter_rows(values_only=True))
                if not rows:
                    raise ValueError("Excel file is empty.")

                header = [str(h).strip().lower() if h is not None else "" for h in rows[0]]

                def find_xl_col(aliases: List[str]) -> int:
                    for alias in aliases:
                        if alias in header:
                            return header.index(alias)
                    return -1

                idx_date = find_xl_col(["date", "ngày", "ngay"])
                idx_cat = find_xl_col(["category", "nhóm", "nhom"])
                idx_asset = find_xl_col(["asset", "tên tài sản", "ten tai san", "tên"])
                idx_action = find_xl_col(["action", "hành động", "hanh dong", "gd"])
                idx_qty = find_xl_col(["quantity", "số lượng", "so luong", "sl"])
                idx_unit = find_xl_col(["unit", "đơn vị", "don vi", "đv"])
                idx_price = find_xl_col(["unitprice", "unit_price", "đơn giá", "don gia", "giá"])
                idx_total = find_xl_col(["total", "tổng tiền", "tong tien", "tổng"])
                idx_note = find_xl_col(["note", "ghi chú", "ghi chu"])

                for row in rows[1:]:
                    if not row or all(v is None for v in row):
                        continue
                    date_val = row[idx_date]
                    if isinstance(date_val, (datetime.date, datetime.datetime)):
                        date_str = date_val.strftime("%Y-%m-%d")
                    else:
                        date_str = str(date_val).strip() if date_val else ""

                    qty = float(str(row[idx_qty]).replace(",", "").strip())
                    price = float(str(row[idx_price]).replace(",", "").strip())
                    total = (
                        float(str(row[idx_total]).replace(",", "").strip())
                        if idx_total != -1 and row[idx_total] is not None
                        else qty * price
                    )

                    imported_records.append({
                        "date": date_str,
                        "category": normalize_category(str(row[idx_cat]).strip() if idx_cat != -1 and row[idx_cat] else "Gold"),
                        "asset": str(row[idx_asset]).strip(),
                        "action": normalize_action(str(row[idx_action]).strip() if idx_action != -1 and row[idx_action] else "Buy"),
                        "quantity": qty,
                        "unit": normalize_unit(str(row[idx_unit]).strip() if idx_unit != -1 and row[idx_unit] else "Tael"),
                        "unit_price": price,
                        "total": total,
                        "note": str(row[idx_note]).strip() if idx_note != -1 and row[idx_note] else "",
                    })

            if not imported_records:
                messagebox.showwarning("Notice", "No valid transaction rows found in file.")
                return

            choice = messagebox.askyesnocancel(
                "Import Mode",
                f"Found {len(imported_records)} transactions.\n\n"
                "- Click [Yes] to MERGE with current database.\n"
                "- Click [No] to REPLACE entire current database.\n"
                "- Click [Cancel] to abort.",
                default=messagebox.YES,
            )
            if choice is None:
                return

            if choice:  # Merge
                for rec in imported_records:
                    rec["id"] = next_id(self.main_app.db)
                    self.main_app.db.append(rec)
                action_text = "Merged"
            else:  # Replace
                for idx, rec in enumerate(imported_records, start=1):
                    rec["id"] = idx
                self.main_app.db = imported_records
                action_text = "Replaced with"

            save_db(self.main_app.db)
            self.refresh_table()
            self.main_app.refresh_portfolio_view()
            messagebox.showinfo("Success", f"{action_text} {len(imported_records)} records successfully.")
            self.main_app.var_status.set(f"Imported {len(imported_records)} transactions.")

        except Exception as e:
            messagebox.showerror("Import Error", f"Failed to import data: {e}")
